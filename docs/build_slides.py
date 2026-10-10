"""Build the 5-slide presentation (docs/presentacion.pptx)."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
SHOTS = ROOT / "docs" / "capturas"
m = json.loads((RES / "metrics.json").read_text(encoding="utf-8"))
tm, tb = m["test_model"], m["test_baseline"]

INK = RGBColor(0x1D, 0x27, 0x31)
ACC = RGBColor(0x1F, 0x6F, 0x8B)
MUT = RGBColor(0x5F, 0x6B, 0x76)

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
BLANK = prs.slide_layouts[6]


def num(x, d=2):
    return f"{x:.{d}f}".replace(".", ",")


def text(slide, x, y, w, h, lines, size=18, color=INK, bold_first=False):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(6)
    return box


def title(slide, t, subtitle=None):
    text(slide, 1.5, 0.9, 31, 1.6, [t], size=28, color=INK, bold_first=True)
    bar = slide.shapes.add_shape(1, Cm(1.5), Cm(2.6), Cm(3), Cm(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACC
    bar.line.fill.background()
    if subtitle:
        text(slide, 1.5, 2.9, 31, 1, [subtitle], size=14, color=MUT)


def table(slide, x, y, w, rows, col_w, size=14):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(0.9 * len(rows)))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Cm(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = t.cell(i, j)
            cell.text = str(v)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(size)
            para.font.bold = i == 0
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACC if i == 0 else (RGBColor(0xF2, 0xF5, 0xF7) if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# 1. Problem and solution
s = prs.slides.add_slide(BLANK)
title(s, "Predicción del consumo diario de insumos", "Proyecto Capstone · Nida Hugo Loza Guzmán")
text(s, 1.5, 4.4, 15, 12, [
    "Problema",
    "En la cafetería los insumos se compran a ojo: a veces falta leche o queso y otras veces se vence la crema o el tomate.",
    "Solución",
    "Predecir cuánto se consume cada día de cada insumo y, con eso, decir cuántos días alcanza el stock y cuánto pedir.",
], size=17)
table(s, 17.5, 4.6, 15, [
    ["Elemento", "Definición"],
    ["Unidad", "Un insumo en una fecha"],
    ["Objetivo", "Consumo del día (salidas del kardex)"],
    ["Línea base", "Mismo día de la semana anterior"],
    ["Modelo", "LightGBM: consumo pasado, stock, calendario"],
    ["Métrica", "MASE (< 1 supera a la base)"],
    ["Validación", "División por fecha"],
], [4, 11], size=13)
notes(s, "Presentar el problema: compras a ojo, faltantes y vencimientos. La idea es usar el kardex que ya existe.")

# 2. Data and model
s = prs.slides.add_slide(BLANK)
title(s, "Dataset y modelo", "Datos reales del sistema de gestión de una cafetería de Tarija (PostgreSQL)")
text(s, 1.5, 4.2, 15.5, 13, [
    "• 374 404 movimientos de kardex, oct-2021 a oct-2026.",
    "• 14 insumos × 1 829 días = 25 606 filas.",
    "• Consumo = producción + venta − devoluciones. Mermas y vencimientos aparte.",
    "• Se excluyen clientes, usuarios y ventas: solo se comparte el agregado diario.",
    "• Variables: rezagos 1–28 días, medias móviles, stock del día anterior, día de la semana, mes, feriados de Bolivia y Tarija.",
    "• Un LightGBM global, cantidades normalizadas por insumo. Pérdida Tweedie elegida en validación.",
    "• Entrenamiento 2021–jun 2025 · Validación jul–dic 2025 · Prueba ene–oct 2026.",
], size=15)
s.shapes.add_picture(str(RES / "importancia_variables.png"), Cm(18), Cm(4), width=Cm(14.5))
notes(s, "Remarcar la división temporal y que la prueba se usó una sola vez.")

# 3. Results
s = prs.slides.add_slide(BLANK)
title(s, "Resultado frente a la línea base", "Conjunto de prueba: 1 de enero a 3 de octubre de 2026")
table(s, 1.5, 4.2, 14.5, [
    ["Métrica (prom. 14 insumos)", "LightGBM", "Base"],
    ["MASE", num(tm["mase_mean"], 3), num(tb["mase_mean"], 3)],
    ["MASE mediana", num(tm["mase_median"], 3), num(tb["mase_median"], 3)],
    ["WAPE diario", f"{num(tm['wape_mean'] * 100, 1)} %", f"{num(tb['wape_mean'] * 100, 1)} %"],
    ["WAPE semanal", f"{num(tm['weekly_wape_mean'] * 100, 1)} %", f"{num(tb['weekly_wape_mean'] * 100, 1)} %"],
    ["Sesgo", f"{num(tm['bias_mean'] * 100, 1)} %", f"{num(tb['bias_mean'] * 100, 1)} %"],
    ["Insumos donde gana", f"{m['items_where_model_wins']} de 14", "–"],
], [7.5, 3.5, 3.5], size=14)
text(s, 1.5, 11.2, 14.5, 6, [
    "• Gana en los 14 insumos; baja el error entre 25 % y 34 %.",
    "• Crema y tomate: MASE > 1 en ambos porque recién se usan desde 2024.",
    "• Harina y huevo (uso por lotes) son el caso difícil.",
], size=14)
s.shapes.add_picture(str(RES / "mase_por_insumo.png"), Cm(16.8), Cm(4), width=Cm(16))
notes(s, "Leer la MASE promedio 0,93 contra 1,32. Explicar por qué crema y tomate superan 1.")

# 4. Demo
s = prs.slides.add_slide(BLANK)
title(s, "Demostración del sistema", "Frontend HTML → API FastAPI → modelo LightGBM (model.txt)")
s.shapes.add_picture(str(SHOTS / "01_caso_valido_leche.png"), Cm(1.5), Cm(4), width=Cm(19))
text(s, 21.3, 4.2, 11.5, 13, [
    "Flujo de la demo",
    "1. Elegir insumo y cargar historial del kardex.",
    "2. Ingresar stock, horizonte y días de entrega.",
    "3. El backend valida, arma variables y ejecuta el modelo.",
    "4. Se ven pronóstico, cobertura y cantidad sugerida.",
    "Caso límite",
    "Historial corto o con días faltantes: el backend responde 422 con el motivo.",
    "Harina con stock 0: avisa pedir hoy.",
], size=14, bold_first=True)
notes(s, "Demo en vivo: http://127.0.0.1:8000. Leche entera, stock 20000, 7 días, entrega 2. Luego borrar líneas del historial para mostrar el error. Respaldo: capturas en docs/capturas.")

# 5. Conclusion
s = prs.slides.add_slide(BLANK)
title(s, "Conclusiones y limitaciones")
text(s, 1.5, 3.6, 15.5, 13, [
    "Conclusiones",
    f"• MASE {num(tm['mase_mean'], 3)} contra {num(tb['mase_mean'], 3)}: el modelo supera a la línea base en todos los insumos.",
    "• Evaluación honesta: división temporal, prueba usada una vez, sin fugas.",
    "• Sistema completo: datos, modelo, API, interfaz y pruebas (10 pasan).",
    "• En el total semanal la ganancia es menor (18,9 % vs 20,9 % WAPE).",
], size=15, bold_first=True)
text(s, 17.5, 3.6, 15, 10, [
    "Limitaciones",
    "• Entrenado a 1 día; a 7–14 días es recursivo.",
    "• Stock = envases abiertos del kardex, no almacén.",
    "• Sin promociones, eventos ni clima.",
    "• Subestima ~10 %; se compensa con stock de seguridad.",
    "• Una sola cafetería.",
], size=15, bold_first=True)
text(s, 1.5, 15.6, 31, 2, ["Repositorio: github.com/hloza/prediccion-insumos-kardex · tag v1.0"], size=15, color=ACC)
notes(s, "Cerrar con el repositorio y el tag v1.0.")

prs.save(ROOT / "docs" / "presentacion.pptx")
print("ok")
